#!/usr/bin/env python3
"""
Generates HouseholdApp/Resources/MealPlanTemplate.xlsx — the spreadsheet a
household member fills out (Excel, Google Sheets, or Numbers) and then
imports from the Meals tab.

Re-run after changing columns:  python3 scripts/make-meal-template.py
Keep this in sync with MealPlanImporter.swift (column aliases) and the
README's "Plan from a spreadsheet" section.

Error-proofing built into the file itself:
  • Every sheet is protected: header rows and the Instructions sheet are
    locked; only the data cells can be edited.
  • Workbook structure is locked so sheets can't be renamed or deleted.
  • Dropdowns for Meal type and Packing section; dates validated as dates.
  • Header cells carry comments explaining required vs optional.
The importer is also tolerant (case-insensitive headers, many date formats,
blank rows ignored) so mild deviations still work.
"""

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.protection import SheetProtection
import os

OUT = os.path.join(os.path.dirname(__file__), "..", "HouseholdApp", "Resources", "MealPlanTemplate.xlsx")
DATA_ROWS = 200          # how many editable rows to prepare per sheet
PROTECT_PASSWORD = None  # no password: protection is a guardrail, not a lock

MEAL_TYPES = ["Breakfast", "Brunch", "Lunch", "Dinner", "Snack", "Dessert"]
SECTIONS   = ["Clothing", "Food", "Toiletries", "Electronics", "Documents", "Other"]

REQ_FILL  = PatternFill("solid", fgColor="FDE9D9")   # soft orange = required
OPT_FILL  = PatternFill("solid", fgColor="EAF1FB")   # soft blue   = optional
HEAD_FONT = Font(bold=True)
THIN      = Side(style="thin", color="BBBBBB")
BORDER    = Border(bottom=THIN)

# (header, required, width, comment)
MEAL_COLS = [
    ("Date",               True,  14, "REQUIRED. The day the meal is planned for.\nType a date (e.g. 2026-09-14 or Sep 14). Excel/Sheets date cells work too."),
    ("Meal",               True,  12, "REQUIRED. Pick from the dropdown: Breakfast, Brunch, Lunch, Dinner, Snack, Dessert."),
    ("Meal Name",          False, 26, "Optional dish name, e.g. \"Spaghetti Bolognese\". Leave blank to just show the meal type."),
    ("Cook",               False, 14, "Optional. A household member's name exactly as it appears in the app. Blank = everyone."),
    ("Ingredients (have)", False, 34, "Optional. Ingredients you already have. Separate with commas: eggs, milk, butter"),
    ("Ingredients to Buy", False, 34, "Optional. Ingredients you need to buy — each becomes a shopping item linked to this meal. Separate with commas."),
    ("Trip Name",          False, 20, "Optional. Ties this meal to a trip. If the trip doesn't exist yet it's created for you (add dates on the Trips sheet, or it spans the meals' dates). All ingredients go on the trip's packing list under Food."),
    ("Recipe Link",        False, 30, "Optional. A web link to the recipe (https://…)."),
    ("Instructions",       False, 40, "Optional. Cooking steps. Line breaks are fine (Alt+Enter in Excel, Ctrl+Enter in Sheets)."),
    ("Notes",              False, 30, "Optional. Anything else."),
]

TRIP_COLS = [
    ("Trip Name",  True,  24, "REQUIRED. Must match the Trip Name used on the Meals and Packing sheets (not case-sensitive)."),
    ("Start Date", True,  14, "REQUIRED. First day of the trip."),
    ("End Date",   True,  14, "REQUIRED. Last day of the trip. Same as Start Date for a single day."),
    ("Notes",      False, 36, "Optional."),
]

PACK_COLS = [
    ("Trip Name", True,  24, "REQUIRED. The trip this item belongs to. Must exist in the app, on the Trips sheet, or be named on a meal."),
    ("Item",      True,  30, "REQUIRED. What to pack, e.g. \"Rain jacket\"."),
    ("Section",   False, 16, "Optional. Pick from the dropdown. Blank = Other."),
]


def write_sheet(ws, cols, freeze="A2"):
    """Header row + formatting + protection with editable data cells."""
    for idx, (title, required, width, comment) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=idx, value=f"{title} *" if required else title)
        cell.font = HEAD_FONT
        cell.fill = REQ_FILL if required else OPT_FILL
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.comment = Comment(comment, "HouseholdApp")
        cell.comment.width = 320
        cell.comment.height = 110
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = freeze

    # Unlock every data cell; the header row stays locked by default.
    for r in range(2, DATA_ROWS + 2):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=r, column=c)
            cell.protection = Protection(locked=False)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.protection = SheetProtection(
        sheet=True, formatCells=False, formatColumns=False, formatRows=False,
        insertRows=False, deleteRows=False, sort=False, autoFilter=False,
        selectLockedCells=False, selectUnlockedCells=False,
    )
    if PROTECT_PASSWORD:
        ws.protection.password = PROTECT_PASSWORD


def add_list_validation(ws, col_letter, options, title):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.error = f"Pick a {title} from the dropdown."
    dv.errorTitle = f"Invalid {title}"
    dv.prompt = f"Choose a {title}"
    dv.promptTitle = title
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{DATA_ROWS + 1}")


def add_date_validation(ws, col_letter):
    dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True)
    dv.error = "Enter a date, e.g. 2026-09-14."
    dv.errorTitle = "Invalid date"
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{DATA_ROWS + 1}")
    for r in range(2, DATA_ROWS + 2):
        ws[f"{col_letter}{r}"].number_format = "yyyy-mm-dd"


def instructions_sheet(ws):
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 90
    rows = [
        ("HouseholdApp — Meal Plan Template", None),
        ("", None),
        ("How to use", "1. Fill in the Meals sheet (one row per meal). Trips and Packing are optional.\n"
                       "2. Save the file as .xlsx (Google Sheets: File → Download → Microsoft Excel).\n"
                       "3. In the app: Meals tab → spreadsheet button → Import. You'll see a preview and can fix anything before it's added."),
        ("", None),
        ("Colour key", "Orange header = required.  Blue header = optional.  Hover a header for details."),
        ("Dates", "Type them any common way: 2026-09-14, Sep 14, 9/14/2026, or use the cell as a real date. Past dates are allowed but not recommended."),
        ("Lists in one cell", "Ingredients are comma-separated:  eggs, milk, butter"),
        ("Names", "Cook must match a household member's name in the app (not case-sensitive). Blank = everyone."),
        ("Trips", "Naming a trip on a meal links it. If the trip isn't in the app or on the Trips sheet, it's created spanning the dates of its meals."),
        ("Duplicates", "Meals that already exist (same day, type and name) and shopping/packing items already on a list are skipped, so re-importing the same file is safe."),
        ("Protection", "Headers and sheet names are locked so the import always works. Only the white cells are editable. Don't unprotect unless you know what you're changing."),
        ("", None),
        ("EXAMPLE — Meals", None),
        ("Date | Meal | Meal Name | Cook | Ingredients (have) | Ingredients to Buy | Trip Name",
         "2026-09-14 | Dinner | Spaghetti Bolognese | Jordan | pasta, garlic | ground beef, tomatoes | \n"
         "2026-09-20 | Breakfast | Pancakes | | flour, eggs | maple syrup | Cottage weekend"),
        ("", None),
        ("EXAMPLE — Trips", None),
        ("Trip Name | Start Date | End Date", "Cottage weekend | 2026-09-19 | 2026-09-21"),
        ("", None),
        ("EXAMPLE — Packing", None),
        ("Trip Name | Item | Section", "Cottage weekend | Rain jacket | Clothing\nCottage weekend | Sunscreen | Toiletries"),
    ]
    for r, (a, b) in enumerate(rows, start=1):
        ca = ws.cell(row=r, column=1, value=a)
        ca.alignment = Alignment(vertical="top", wrap_text=True)
        if r == 1:
            ca.font = Font(bold=True, size=14)
        elif a and not b:
            ca.font = Font(bold=True)
        elif a:
            ca.font = Font(bold=True, color="444444")
        if b:
            cb = ws.cell(row=r, column=2, value=b)
            cb.alignment = Alignment(vertical="top", wrap_text=True)
    ws.protection = SheetProtection(sheet=True)


def main():
    wb = Workbook()
    ws_i = wb.active
    ws_i.title = "Instructions"
    instructions_sheet(ws_i)

    ws_m = wb.create_sheet("Meals")
    write_sheet(ws_m, MEAL_COLS)
    add_date_validation(ws_m, "A")
    add_list_validation(ws_m, "B", MEAL_TYPES, "Meal")

    ws_t = wb.create_sheet("Trips")
    write_sheet(ws_t, TRIP_COLS)
    add_date_validation(ws_t, "B")
    add_date_validation(ws_t, "C")

    ws_p = wb.create_sheet("Packing")
    write_sheet(ws_p, PACK_COLS)
    add_list_validation(ws_p, "C", SECTIONS, "Section")

    # Open on Meals; lock sheet structure (no rename/delete/add).
    wb.active = 1
    wb.security = wb.security or None
    from openpyxl.workbook.protection import WorkbookProtection
    wb.security = WorkbookProtection(workbookPassword=None, lockStructure=True)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("wrote", os.path.normpath(OUT))


if __name__ == "__main__":
    main()
